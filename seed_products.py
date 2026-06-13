#!/usr/bin/env python3
"""
seed_products.py — Load ElmShelf.xlsx product data into PostgreSQL.

Run from the E-Commerce-Backend directory:
    python seed_products.py
    python seed_products.py "path/to/ElmShelf.xlsx"   # custom path
"""
import os
import re
import sys
import unicodedata

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import SessionLocal
from app.models.category import Category
from app.models.inventory import Inventory
from app.models.product import (
    Product, ProductType, ProductVariant, ProductVariantGroup,
)


# ── helpers ───────────────────────────────────────────────────────────────────

def slugify(s: str) -> str:
    s = str(s).lower().strip()
    s = unicodedata.normalize("NFKD", s)
    s = re.sub(r"[^\w\s-]", "", s, flags=re.ASCII)
    s = re.sub(r"[\s_]+", "-", s)
    return re.sub(r"-+", "-", s).strip("-")


def unique_slug(base: str, seen: set) -> str:
    slug, i = base, 1
    while slug in seen:
        slug = f"{base}-{i}"
        i += 1
    seen.add(slug)
    return slug


def parse_price(raw) -> tuple:
    """Return (min_price: float, max_price: float | None)."""
    if not raw:
        return None, None
    s = re.sub(r"[££Â$€,]", "", str(raw)).strip()
    if not s or s in ("-", "--", "N/A", "POA", "TBC", "0"):
        return None, None
    m = re.search(r"(\d[\d.]*)\s*[–\-]+\s*(\d[\d.]*)", s)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        return (lo, hi) if lo <= hi else (hi, lo)
    m = re.search(r"(\d[\d.]*)", s)
    return (float(m.group(1)), None) if m else (None, None)


def clean(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ("N/A",) or re.fullmatch(r"-+", s):
        return None
    return s


def split_colours(val: str) -> list:
    val = re.sub(r"\s+or\s+", ",", val, flags=re.IGNORECASE)
    parts = [p.strip() for p in re.split(r"[,/]", val)]
    return [p for p in parts if p and p not in ("-", "--")]


def text_to_html(raw: str | None) -> str | None:
    """Convert plain text with newlines into TipTap-compatible HTML paragraphs."""
    if not raw:
        return None
    paragraphs = re.split(r"\n{2,}", raw.strip())
    parts = []
    for para in paragraphs:
        para = para.strip()
        if para:
            inner = para.replace("\n", "<br>")
            parts.append(f"<p>{inner}</p>")
    return "".join(parts) if parts else None


def specs_to_html(specs: dict) -> str | None:
    """Build an HTML bullet list from product spec key-value pairs."""
    if not specs:
        return None
    labels = {
        "price_range_text": "Price Range",
        "depth_mm":         "Depth (mm)",
        "width_mm":         "Width (mm)",
        "height_mm":        "Height (mm)",
        "weight_kg":        "Weight (kg)",
        "finish":           "Finish",
        "size":             "Size",
        "cable_tidy":       "Cable Tidy",
        "model":            "Model",
        "colour":           "Colour",
    }
    items = ""
    for key, val in specs.items():
        label = labels.get(key, key.replace("_", " ").title())
        val_html = str(val).strip().replace("\n", "<br>")
        items += f"<li><strong>{label}:</strong> {val_html}</li>"
    return f"<ul>{items}</ul>" if items else None


# ── column mapping ────────────────────────────────────────────────────────────

SKIP_SHEETS = {"Sheet3"}

COL_ALIAS = {
    "product name":            "name",
    "category":                "category",
    "sub category":            "sub_category",
    "price range":             "price_range",
    "colour":                  "colour",
    "color":                   "colour",
    "depth(mm)":               "depth_mm",
    "length(mm)":              "depth_mm",
    "width(mm)":               "width_mm",
    "height(mm)":              "height_mm",
    "weight(kg)":              "weight_kg",
    "finish":                  "finish",
    "size":                    "size",
    "cable tidy":              "cable_tidy",
    "cabel tidy":              "cable_tidy",   # tolerate typo in sheet
    "model":                   "model",
    "main note":               "main_note",
    "description":             "description",
    "key features":            "key_features",
    "what's included":         "whats_included",
    "important notes":         "important_notes",
    "additional information":  "additional_info",
    "additional information ": "additional_info",
}


def find_header(ws):
    """Return (header_row_number, canonical_col_map {key: col_index}).
    Scans first 6 rows for a row containing 'product name'.
    """
    for row_num, row in enumerate(ws.iter_rows(max_row=6), start=1):
        lower = [str(c.value).strip().lower() if c.value else "" for c in row]
        if "product name" in lower:
            col_map: dict = {}
            for idx, lv in enumerate(lower):
                canon = COL_ALIAS.get(lv, lv) if lv else None
                if canon and canon not in col_map:
                    col_map[canon] = idx
            return row_num, col_map
    return None, {}


def vget(vals: list, col_map: dict, key: str):
    idx = col_map.get(key)
    if idx is None or idx >= len(vals):
        return None
    return clean(vals[idx])


# ── seeder ────────────────────────────────────────────────────────────────────

def seed(xlsx_path: str):
    db: Session = SessionLocal()
    try:
        # 1. Clear existing product & category data
        print("Clearing existing product data...")
        for tbl in [
            "stock_movements",
            "inventory",
            "reviews",
            "order_items",
            "cart_items",
            "product_variants",
            "product_variant_groups",
            "product_images",
            "product_categories",
            "products",
            "categories",
        ]:
            db.execute(text(f"DELETE FROM {tbl}"))
        db.commit()
        print("  Done.\n")

        # 2. Open workbook
        print(f"Reading: {xlsx_path}")
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)

        seen_slugs: set = set()
        cat_cache: dict = {}
        sku_idx = 1
        total = 0

        def get_or_create_cat(name: str, parent_id=None) -> Category:
            key = f"{parent_id}::{name}"
            if key in cat_cache:
                return cat_cache[key]
            sl = unique_slug(slugify(name), seen_slugs)
            cat = Category(name=name, slug=sl, parent_id=parent_id, is_active=True)
            db.add(cat)
            db.flush()
            cat_cache[key] = cat
            return cat

        # 3. Process each sheet
        for sheet_name in wb.sheetnames:
            if sheet_name in SKIP_SHEETS:
                continue

            ws = wb[sheet_name]
            hdr_row, col_map = find_header(ws)
            if not hdr_row:
                print(f"  [SKIP] {sheet_name} — header row not found")
                continue

            sheet_count = 0
            skipped = 0

            for row in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
                vals = list(row)
                name = vget(vals, col_map, "name")
                if not name:
                    skipped += 1
                    continue

                # --- Categories ---
                cat_name    = vget(vals, col_map, "category") or sheet_name
                subcat_name = vget(vals, col_map, "sub_category")
                parent_cat  = get_or_create_cat(cat_name)
                leaf_cat    = (
                    get_or_create_cat(subcat_name, parent_cat.id)
                    if subcat_name else parent_cat
                )

                # --- Price ---
                price_raw = vget(vals, col_map, "price_range")
                min_price, max_price = parse_price(price_raw)
                if min_price is None:
                    min_price = 0.0

                # --- Colour → variant detection ---
                colour_raw  = vget(vals, col_map, "colour")
                colours     = split_colours(colour_raw) if colour_raw else []
                is_variable = len(colours) > 1

                # --- Rich-text content fields (plain text → HTML) ---
                main_note       = text_to_html(vget(vals, col_map, "main_note"))
                description     = text_to_html(vget(vals, col_map, "description"))
                key_features    = text_to_html(vget(vals, col_map, "key_features"))
                whats_included  = text_to_html(vget(vals, col_map, "whats_included"))
                important_notes = text_to_html(vget(vals, col_map, "important_notes"))

                # --- Specs → HTML list for additional_information ---
                specs: dict = {}
                if price_raw and max_price is not None:
                    specs["price_range_text"] = price_raw
                for fld in [
                    "depth_mm", "width_mm", "height_mm", "weight_kg",
                    "finish", "size", "cable_tidy", "model",
                ]:
                    v = vget(vals, col_map, fld)
                    if v:
                        specs[fld] = v
                if colour_raw:
                    specs["colour"] = colour_raw

                specs_html     = specs_to_html(specs)
                add_info_html  = text_to_html(vget(vals, col_map, "additional_info"))

                if specs_html and add_info_html:
                    additional_information = specs_html + add_info_html
                else:
                    additional_information = specs_html or add_info_html

                # --- Slug & SKU ---
                slug = unique_slug(slugify(name), seen_slugs)
                sku  = f"ELM-{sku_idx:05d}"
                sku_idx += 1

                # --- Product ---
                product = Product(
                    name=name,
                    slug=slug,
                    description=description,
                    main_note=main_note,
                    key_features=key_features,
                    whats_included=whats_included,
                    important_notes=important_notes,
                    additional_information=additional_information,
                    price=min_price,
                    sale_price=max_price,
                    stock_quantity=0,
                    sku=sku,
                    product_type=ProductType.variable if is_variable else ProductType.simple,
                    is_active=True,
                    is_featured=False,
                )
                db.add(product)
                db.flush()

                product.categories.append(leaf_cat)
                if subcat_name and leaf_cat.id != parent_cat.id:
                    product.categories.append(parent_cat)
                db.add(Inventory(product_id=product.id, on_hand=0, reserved=0))

                # --- Colour variant group ---
                if is_variable:
                    grp = ProductVariantGroup(product_id=product.id, attribute="Color")
                    db.add(grp)
                    db.flush()
                    for colour in colours:
                        db.add(ProductVariant(
                            group_id=grp.id,
                            value=colour,
                            price_modifier=0,
                            stock_quantity=0,
                        ))

                sheet_count += 1
                total += 1

                if total % 100 == 0:
                    db.commit()
                    print(f"    {total} products committed...")

            db.commit()
            print(f"  [{sheet_name}] {sheet_count} products loaded"
                  + (f"  ({skipped} empty rows skipped)" if skipped else ""))

        print(f"\n{'='*50}")
        print(f"  Total products : {total}")
        print(f"  Total categories: {len(cat_cache)}")
        print(f"{'='*50}")

    except Exception as e:
        db.rollback()
        print(f"\nERROR — rolled back: {e}", file=sys.stderr)
        raise
    finally:
        db.close()


if __name__ == "__main__":
    xlsx_path = (
        sys.argv[1]
        if len(sys.argv) > 1
        else os.path.normpath(
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "ElmShelf.xlsx")
        )
    )
    if not os.path.exists(xlsx_path):
        print(f"ERROR: File not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)
    seed(xlsx_path)
